import glob
import os
import re
import uuid

import chromadb
from chromadb.utils import embedding_functions
from docx import Document
from pypdf import PdfReader
from tqdm import tqdm


DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "data")
DB_DIR = os.path.join(os.path.dirname(__file__), "..", "chroma_db")
COLLECTION_NAME = "documents"

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}


def read_text_file(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def read_pdf_file(path: str) -> str:
    reader = PdfReader(path)

    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""

        if text.strip():
            pages.append(text)

    return "\n\n".join(pages)


def read_docx_file(path: str) -> str:
    document = Document(path)

    parts = []

    # Extract normal paragraphs.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Extract table contents as well.
    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                parts.append(" | ".join(cells))

    return "\n\n".join(parts)


def read_document(path: str) -> str:
    extension = os.path.splitext(path)[1].lower()

    if extension == ".pdf":
        return read_pdf_file(path)

    if extension == ".docx":
        return read_docx_file(path)

    if extension == ".txt":
        return read_text_file(path)

    raise ValueError(f"Unsupported file type: {extension}")


def clean_text(text: str) -> str:
    """Clean common extraction artifacts while preserving structure."""

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Normalize spaces and tabs.
    text = re.sub(r"[ \t]+", " ", text)

    # Reduce excessive blank lines.
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Join words broken across lines with a hyphen.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)

    # Join normal wrapped lines while preserving paragraph breaks.
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)

    return text.strip()


def split_into_paragraphs(text: str) -> list[str]:
    """Split text using paragraph boundaries first."""

    paragraphs = re.split(r"\n\s*\n", text)

    return [
        paragraph.strip()
        for paragraph in paragraphs
        if paragraph.strip()
    ]


def chunk_text(
    text: str,
    chunk_size: int = CHUNK_SIZE,
    overlap: int = CHUNK_OVERLAP,
) -> list[str]:

    paragraphs = split_into_paragraphs(text)

    chunks = []
    current = ""

    for paragraph in paragraphs:

        candidate = (
            f"{current}\n\n{paragraph}"
            if current
            else paragraph
        )

        if len(candidate) <= chunk_size:
            current = candidate
            continue

        if current:
            chunks.append(current.strip())

        # Handle unusually large paragraphs.
        if len(paragraph) > chunk_size:
            start = 0

            while start < len(paragraph):
                end = start + chunk_size

                piece = paragraph[start:end].strip()

                if piece:
                    chunks.append(piece)

                start += chunk_size - overlap

            current = ""

        else:
            current = paragraph

    if current:
        chunks.append(current.strip())

    return chunks


def get_semester_and_course(path: str) -> tuple[str, str]:
    """
    Extract semester and course from the folder structure.

    Expected structure:

        data/
            Rag Notes/
                1_Notes/
                    Python/
                        notes.pdf

    Returns:

        semester = "1"
        course = "Python"
    """

    relative_path = os.path.relpath(path, DATA_DIR)

    parts = relative_path.split(os.sep)

    semester = "Unknown"
    course = "Unknown"

    semester_index = None

    for index, part in enumerate(parts):
        match = re.fullmatch(r"(\d+)_Notes", part)

        if match:
            semester = match.group(1)
            semester_index = index
            break

    if semester_index is not None:
        course_index = semester_index + 1

        if course_index < len(parts) - 1:
            course = parts[course_index]

    return semester, course


def discover_documents() -> list[str]:
    """Recursively discover all supported documents."""

    paths = []

    for root, _, files in os.walk(DATA_DIR):

        for filename in files:

            extension = os.path.splitext(filename)[1].lower()

            if extension in SUPPORTED_EXTENSIONS:
                paths.append(os.path.join(root, filename))

    return sorted(paths)


def load_documents():
    paths = discover_documents()

    if not paths:
        raise FileNotFoundError(
            f"No supported documents found in {DATA_DIR}.\n"
            f"Supported types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    docs = []

    for path in paths:

        filename = os.path.basename(path)

        relative_path = os.path.relpath(path, DATA_DIR)

        semester, course = get_semester_and_course(path)

        print(f"Reading: {relative_path}")
        print(f"  Semester: {semester}")
        print(f"  Course: {course}")

        try:
            text = read_document(path)
            text = clean_text(text)

            if not text:
                print("  WARNING: No text extracted. Skipping.")
                continue

            chunks = chunk_text(text)

            print(f"  Created {len(chunks)} chunks.")

            for index, chunk in enumerate(chunks):

                docs.append(
                    {
                        "id": str(uuid.uuid4()),
                        "text": chunk,
                        "source": filename,
                        "relative_path": relative_path,
                        "semester": semester,
                        "course": course,
                        "file_type": os.path.splitext(filename)[1].lower(),
                        "chunk_index": index,
                    }
                )

        except Exception as error:
            print(f"  ERROR: {error}")
            print("  Skipping this file.")

    return docs


def main():

    print("=" * 60)
    print("RAG DOCUMENT INGESTION")
    print("=" * 60)

    print(f"\nDocument root: {os.path.abspath(DATA_DIR)}")

    docs = load_documents()

    print(f"\nLoaded {len(docs)} total chunks.")

    if not docs:
        raise RuntimeError("No documents were successfully processed.")

    embedding_fn = (
        embedding_functions.SentenceTransformerEmbeddingFunction(
            model_name="all-MiniLM-L6-v2"
        )
    )

    client = chromadb.PersistentClient(path=DB_DIR)

    try:
        client.delete_collection(COLLECTION_NAME)
        print("Deleted existing collection.")

    except Exception:
        pass

    collection = client.create_collection(
        name=COLLECTION_NAME,
        embedding_function=embedding_fn,
    )

    batch_size = 64

    for i in tqdm(
        range(0, len(docs), batch_size),
        desc="Embedding + indexing",
    ):

        batch = docs[i:i + batch_size]

        collection.add(
            ids=[document["id"] for document in batch],

            documents=[
                document["text"]
                for document in batch
            ],

            metadatas=[
                {
                    "source": document["source"],
                    "relative_path": document["relative_path"],
                    "semester": document["semester"],
                    "course": document["course"],
                    "file_type": document["file_type"],
                    "chunk_index": document["chunk_index"],
                }
                for document in batch
            ],
        )

    print("\n" + "=" * 60)
    print("INGESTION COMPLETE")
    print("=" * 60)

    print(f"Indexed chunks : {len(docs)}")
    print(f"Vector store   : {os.path.abspath(DB_DIR)}")
    print(f"Collection     : {COLLECTION_NAME}")


if __name__ == "__main__":
    main()